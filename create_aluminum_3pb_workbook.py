"""Create the editable aluminum input template (requires openpyxl).

Existing workbooks are preserved unless --force is explicitly supplied.
"""
from pathlib import Path
import argparse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment


def create(path):
    wb = Workbook()
    wb.remove(wb.active)
    blue = PatternFill('solid', fgColor='17365D')
    yellow = PatternFill('solid', fgColor='FFF2CC')

    def sheet(name, headers, rows, widths):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = blue
            cell.font = Font(color='FFFFFF', bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical='center')
        ws.row_dimensions[1].height = 30
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name='Calibri', size=11)
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                cell.fill = yellow
                if isinstance(cell.value, float):
                    cell.number_format = '0.########'
        for i, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = width
        if name != 'Instructions':
            tab = Table(displayName='Input_' + name, ref=ws.dimensions)
            tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
            ws.add_table(tab)
        return ws

    def flag(ws, col):
        dv = DataValidation(type='list', formula1='"0,1"', allow_blank=False)
        dv.errorTitle = 'Use 0 or 1'
        dv.error = '0 disables this item; 1 enables it.'
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}10000')

    sheet('Instructions', ['Topic', 'Instructions'], [
        ['START', 'Edit this workbook, save it, then run run_aluminum_3pb_excel in MATLAB.'],
        ['Units', 'N, mm, MPa. Coordinates: x across flange, y along length, z upward.'],
        ['Editable cells', 'Yellow cells are inputs. Keep sheet names, column headers and Parameter keys unchanged.'],
        ['Example only', 'ALL material curves are illustrative. Replace them with your alloy/temper and HAZ test data.'],
        ['Geometry', 'height is the overall height. topWidth and bottomWidth are full flange widths.'],
        ['Three-point bending', 'Supports at y=0 and y=length, central downward load at y=length/2. Ideal bearing lines; no contact.'],
        ['Stiffeners', 'One row per transverse plate. Enabled=1 uses the row; 0 ignores it. Side +1 or -1 selects the web side.'],
        ['Position and width', 'Position is absolute y in mm. Width is from the web mid-plane; do not exceed either flange half-width.'],
        ['Changing length', 'Update stiffener positions manually after changing length. Rows do not automatically move.'],
        ['Disable stiffeners', 'Set all Enabled values to 0, or delete all data rows while retaining the headers.'],
        ['HAZ', 'flangeWidth and webWidth apply to longitudinal welds; stiffenerWidth applies to transverse welds. See ALUMINUM_3PB.md.'],
        ['Materials IDs 1-8', '1 web base; 2 web HAZ; 3 top base; 4 top HAZ; 5 bottom base; 6 bottom HAZ; 7 stiffener base; 8 stiffener HAZ.'],
        ['Additional materials', 'Append contiguous IDs 9,10,... for custom Regions. Names must be unique letters/digits/underscores starting with a letter.'],
        ['Plastic data', 'Rows contain MaterialID, TRUE stress (MPa), equivalent PLASTIC strain (dimensionless). First strain=0; subsequent strains strictly increase. Stress must not decrease.'],
        ['Elastic only', 'UsePlastic=0 ignores that material\'s Plastic rows. UsePlastic=1 requires a curve. Preserve the row order of each curve.'],
        ['Regions', 'Enabled box regions override automatic HAZ. Component: 0 all, 1 web, 2 top, 3 bottom, 4 stiffeners. Later rows take priority.'],
        ['Region bounds', 'Blank minimum = minus infinity; blank maximum = plus infinity. Bounds refer to undeformed x/y/z.'],
        ['Output directory', 'A relative outputDir is resolved relative to this workbook. An absolute directory may also be used.'],
        ['Boolean inputs', 'Use numeric 0 or 1 for preview, lateralGuideAtLoad, Enabled and UsePlastic.'],
        ['Formula inputs', 'Prefer plain numbers. Formula results must be recalculated and saved in Excel before MATLAB reads the workbook.'],
        ['MATLAB command', "run_aluminum_3pb_excel('Aluminum_3PB_Input.xlsx')"],
        ['Output', 'Abaqus .inp, mesh/config .mat, element-region .csv and optional preview. The command does not launch Abaqus.'],
    ], [28, 112])
    for i in range(2, 24):
        wb['Instructions'].row_dimensions[i].height = 34
    control = sheet('Control', ['Parameter', 'Value', 'Units', 'Description'], [
        ['name', 'Aluminum_Welded_I_3PB', '-', 'Abaqus job name, letters/digits/underscores, start with a letter'],
        ['outputDir', 'Aluminum_3PB_output', '-', 'Relative to workbook folder, or absolute folder'],
        ['meshSize', 20, 'mm', 'Maximum target mesh interval; HAZ boundaries may produce smaller elements'],
        ['imperfection', 0, 'mm', 'Sinusoidal lateral crookedness amplitude'],
        ['lateralGuideAtLoad', 1, '0/1', 'Enable lateral guide at loading-line center'],
        ['preview', 1, '0/1', 'Show MATLAB material-region plot'],
    ], [28, 34, 14, 83])
    dv = DataValidation(type='list', formula1='"0,1"'); dv.showErrorMessage = True
    control.add_data_validation(dv); dv.add('B6:B7')
    sheet('Geometry', ['Parameter', 'Value', 'Units', 'Description'], [
        ['length', 2400, 'mm', 'Support-to-support span'],
        ['height', 300, 'mm', 'Overall section height'],
        ['topWidth', 150, 'mm', 'Full upper-flange width'],
        ['bottomWidth', 150, 'mm', 'Full lower-flange width'],
        ['topThickness', 10, 'mm', 'Upper-flange thickness'],
        ['bottomThickness', 10, 'mm', 'Lower-flange thickness'],
        ['webThickness', 6, 'mm', 'Web thickness'],
    ], [28, 22, 14, 83])
    sheet('HAZ', ['Parameter', 'Value', 'Units', 'Description'], [
        ['flangeWidth', 20, 'mm', 'Longitudinal weld: outward HAZ width from web face into flange'],
        ['webWidth', 20, 'mm', 'Longitudinal weld: HAZ depth from flange inner face into clear web'],
        ['stiffenerWidth', 15, 'mm', 'Transverse weld HAZ width from plate/web/flange faces'],
    ], [28, 22, 14, 83])
    sheet('Analysis', ['Parameter', 'Value', 'Units', 'Description'], [
        ['displacement', 40, 'mm', 'Positive downward displacement magnitude at central load line'],
        ['initialIncrement', .01, '-', 'Initial fraction of step time'],
        ['minimumIncrement', 1e-7, '-', 'Minimum fraction of step time'],
        ['maximumIncrement', .05, '-', 'Maximum fraction of step time'],
        ['maximumIncrements', 1000, '-', 'Maximum number of increments'],
    ], [28, 22, 14, 83])
    stiff = sheet('Stiffeners', ['Enabled', 'Position', 'Side', 'Width', 'Thickness'],
                  [[1, y, side, 65, 8] for y in (0, 1200, 2400) for side in (-1, 1)], [18]*5)
    flag(stiff, 'A')
    dv = DataValidation(type='list', formula1='"-1,1"'); dv.showErrorMessage = True
    stiff.add_data_validation(dv); dv.add('C2:C10000')
    names = ['WEB_BASE', 'WEB_HAZ', 'TOP_BASE', 'TOP_HAZ', 'BOTTOM_BASE', 'BOTTOM_HAZ', 'STIFFENER_BASE', 'STIFFENER_HAZ']
    mats = sheet('Materials', ['ID', 'Name', 'E', 'Nu', 'UsePlastic'],
                 [[i, name, 70000, .33, 1] for i, name in enumerate(names, 1)], [12, 28, 20, 18, 18])
    flag(mats, 'E')
    mats['B1'].comment = Comment('Example material names and curves: replace with your measured aluminum properties.', 'Input template')
    strains = [0, .005, .02, .05, .1]
    plastic = [[i, stress, strain] for i in range(1, 9)
               for stress, strain in zip([240, 255, 275, 290, 300] if i % 2 else [110, 125, 145, 165, 180], strains)]
    sheet('Plastic', ['MaterialID', 'TrueStress', 'PlasticStrain'], plastic, [20, 24, 24])
    regions = sheet('Regions', ['Enabled', 'Component', 'Xmin', 'Xmax', 'Ymin', 'Ymax', 'Zmin', 'Zmax', 'MaterialID'],
                    [[0, 1, None, None, 900, 1500, 30, 80, 2]], [16]*9)
    flag(regions, 'A')
    wb.active = 0
    wb.save(path)
    # Round-trip assertions check the actual artifact, including its schema.
    loaded = load_workbook(path, data_only=True)
    assert loaded.sheetnames == ['Instructions', 'Control', 'Geometry', 'HAZ', 'Analysis', 'Stiffeners', 'Materials', 'Plastic', 'Regions']
    assert loaded['Materials'].max_row == 9 and loaded['Plastic'].max_row == 41
    assert loaded['Stiffeners'].max_row == 7
    assert loaded['Geometry']['B2'].value == 2400
    assert loaded['Regions']['C2'].value is None
    assert loaded['Analysis']['B4'].value == 1e-7
    print(f'Created and checked {path}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', nargs='?', type=Path, default=Path(__file__).with_name('Aluminum_3PB_Input.xlsx'))
    parser.add_argument('--force', action='store_true')
    args = parser.parse_args()
    if args.output.exists() and not args.force:
        parser.error('Workbook already exists; use another filename or --force to replace it.')
    create(args.output)
